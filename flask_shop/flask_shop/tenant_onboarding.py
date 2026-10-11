# -*- coding: utf-8 -*-
"""Formal tenant onboarding, initial data import, and demo-data cleanup."""
from __future__ import annotations

import csv
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from sqlalchemy import text

from flask_shop import db, models
from flask_shop.utils.rbac import ALL_ACTION_CODES, ROLE_TEMPLATES
from flask_shop.utils.tenant_query import tenant_context


DEFAULT_PAY_GUIDE = (
    "请按店铺公布的收款方式付款，并在备注中填写订单号。"
    "店员核对到账后会确认收款并安排发货。"
)

MENU_L1 = [
    ("用户管理", "users"),
    ("权限管理", "rights"),
    ("商品管理", "goods"),
    ("订单管理", "orders"),
    ("数据统计", "reports"),
    ("客服中心", "service"),
    ("店铺设置", "shop"),
]

MENU_L2 = [
    ("用户列表", "user_list", "用户管理"),
    ("角色列表", "role_list", "权限管理"),
    ("权限分配", "author_list", "权限管理"),
    ("动作权限", "role_actions", "权限管理"),
    ("商品列表", "product_list", "商品管理"),
    ("添加商品", "add_product", "商品管理"),
    ("分类参数", "attribute_list", "商品管理"),
    ("商品分类", "group_list", "商品管理"),
    ("运费模板", "freight", "商品管理"),
    ("订单列表", "order_list", "订单管理"),
    ("数据报表", "data_list", "数据统计"),
    ("客服工作台", "customer_service", "客服中心"),
    ("店铺设置", "shop_settings", "店铺设置"),
]

TEMPLATE_HEADERS = {
    "products": ["name", "price", "stock", "category", "description", "publish"],
    "staff": ["username", "nickname", "role", "password", "phone", "email"],
    "orders": [
        "order_no", "customer_phone", "customer_name", "product_name", "quantity",
        "unit_price", "status", "created_at", "paid_at", "address",
    ],
}

TEMPLATE_EXAMPLES = {
    "products": ["示例商品", 99.00, 20, "默认分类", "请替换为真实商品说明", "no"],
    "staff": ["service01", "客服一号", "客服", "ChangeMe_2026", "13800000000", ""],
    "orders": [
        "HISTORY-001", "13800000001", "历史顾客", "示例商品", 1, 99.00,
        "completed", "2026-01-01 10:00:00", "2026-01-01 10:05:00", "历史订单地址",
    ],
}

VALID_THEMES = {"retail", "tea", "fashion", "digital"}
VALID_ORDER_STATUS = {"pending_pay", "paid", "shipped", "completed", "closed", "refunding", "refunded"}
PAID_STATUS = {"paid", "shipped", "completed", "refunding", "refunded"}
DEMO_CUSTOMER_PHONES = {"13800001001", "13800001002", "13800001003", "13800001004"}


class OnboardingError(ValueError):
    pass


def ensure_base_menus() -> list[models.Menu]:
    parents: dict[str, models.Menu] = {}
    for name, path in MENU_L1:
        menu = models.Menu.query.filter_by(path=path, level=1).first()
        if not menu:
            menu = models.Menu(name=name, path=path, level=1)
            db.session.add(menu)
            db.session.flush()
        parents[name] = menu

    result = list(parents.values())
    for name, path, parent_name in MENU_L2:
        menu = models.Menu.query.filter_by(path=path).first()
        if not menu:
            menu = models.Menu(name=name, path=path, level=2, pid=parents[parent_name].id)
            db.session.add(menu)
            db.session.flush()
        result.append(menu)
    return result


def _validate_slug(slug: str) -> str:
    value = (slug or "").strip().lower()
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{2,62}", value):
        raise OnboardingError("团队标识只能使用 3-63 位小写字母、数字和连字符")
    if value == "demo" or value.startswith("demo-"):
        raise OnboardingError("正式租户不能使用 demo 或 demo- 开头的团队标识")
    return value


def _validate_password(password: str) -> str:
    value = str(password or "")
    if len(value) < 8:
        raise OnboardingError("初始密码至少 8 位")
    if value in {"123456", "password", "admin123"}:
        raise OnboardingError("不能使用演示或常见弱密码")
    if not re.search(r"[A-Za-z]", value) or not re.search(r"\d", value):
        raise OnboardingError("初始密码必须同时包含字母和数字")
    return value


def create_formal_tenant(
    *,
    slug: str,
    name: str,
    admin_name: str,
    admin_password: str,
    admin_nickname: str = "店长",
    contact_phone: str = "",
    plan_code: str = "pro",
    theme: str = "retail",
    pay_guide: str = DEFAULT_PAY_GUIDE,
    freight_fee: Decimal | float | str = Decimal("0"),
    free_shipping_amount: Decimal | float | str | None = None,
) -> dict[str, Any]:
    slug = _validate_slug(slug)
    name = (name or "").strip()
    admin_name = (admin_name or "").strip()
    password = _validate_password(admin_password)
    if not name:
        raise OnboardingError("店铺名称不能为空")
    if not re.fullmatch(r"[A-Za-z0-9_.-]{3,32}", admin_name):
        raise OnboardingError("管理员账号只能使用 3-32 位字母、数字、点、下划线或连字符")
    if theme not in VALID_THEMES:
        raise OnboardingError(f"店铺主题必须是：{', '.join(sorted(VALID_THEMES))}")
    try:
        fee = Decimal(str(freight_fee or 0))
        free_amount = None if free_shipping_amount in (None, "") else Decimal(str(free_shipping_amount))
    except InvalidOperation as exc:
        raise OnboardingError("运费金额格式不正确") from exc
    if fee < 0 or (free_amount is not None and free_amount < 0):
        raise OnboardingError("运费金额不能小于 0")

    with tenant_context(disabled=True):
        if models.Tenant.query.filter_by(slug=slug).first():
            raise OnboardingError(f"团队标识 {slug} 已存在")

        menus = ensure_base_menus()
        tenant = models.Tenant(
            name=name[:64],
            slug=slug,
            status=1,
            contact_phone=(contact_phone or "").strip()[:32] or None,
            plan_code=(plan_code or "pro").strip()[:32],
            pay_guide=(pay_guide or DEFAULT_PAY_GUIDE).strip(),
            data_mode="live",
            theme=theme,
        )
        db.session.add(tenant)
        db.session.flush()

        role = models.Role(
            tenant_id=tenant.id,
            name="超级管理员",
            desc="正式租户全部权限",
            perm_codes=__import__("json").dumps(ALL_ACTION_CODES, ensure_ascii=False),
            is_super=1,
        )
        role.menus = menus
        db.session.add(role)
        db.session.flush()

        admin = models.User(
            tenant_id=tenant.id,
            name=admin_name,
            nick_name=(admin_nickname or "店长").strip()[:32],
            role_id=role.id,
            is_admin=1,
        )
        admin.password = password
        db.session.add(admin)
        db.session.add(models.FreightTemplate(
            tenant_id=tenant.id,
            name="默认运费",
            charge_type="fixed",
            fixed_fee=fee,
            free_amount=free_amount,
            is_default=1,
            status=1,
        ))
        db.session.commit()
        return {
            "tenant_id": tenant.id,
            "slug": tenant.slug,
            "name": tenant.name,
            "data_mode": tenant.data_mode,
            "admin": admin.name,
            "freight_fee": float(fee),
            "free_shipping_amount": float(free_amount) if free_amount is not None else None,
        }


def export_templates(output_dir: str | Path) -> dict[str, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    destination = Path(output_dir).expanduser().resolve()
    destination.mkdir(parents=True, exist_ok=True)
    files: dict[str, str] = {}

    for kind, headers in TEMPLATE_HEADERS.items():
        path = destination / f"{kind}.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
        files[kind] = str(path)

    workbook = Workbook()
    workbook.remove(workbook.active)
    for kind, headers in TEMPLATE_HEADERS.items():
        sheet = workbook.create_sheet(kind)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 36)
            sheet.column_dimensions[column[0].column_letter].width = width
    guide = workbook.create_sheet("field_guide")
    guide.append(["data_sheet", "field", "example"])
    for cell in guide[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for kind, headers in TEMPLATE_HEADERS.items():
        for field, example in zip(headers, TEMPLATE_EXAMPLES[kind]):
            guide.append([kind, field, example])
    guide.freeze_panes = "A2"
    guide.column_dimensions["A"].width = 16
    guide.column_dimensions["B"].width = 24
    guide.column_dimensions["C"].width = 36
    workbook_path = destination / "onboarding.xlsx"
    workbook.save(workbook_path)
    files["workbook"] = str(workbook_path)
    return files


def _rows_from_sheet(sheet) -> list[dict[str, Any]]:
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [str(value or "").strip() for value in values[0]]
    rows = []
    for values_row in values[1:]:
        row = {headers[index]: value for index, value in enumerate(values_row) if index < len(headers) and headers[index]}
        if any(value not in (None, "") for value in row.values()):
            rows.append(row)
    return rows


def load_import_rows(file_path: str | Path, kind: str) -> dict[str, list[dict[str, Any]]]:
    from openpyxl import load_workbook

    path = Path(file_path).expanduser().resolve()
    if not path.is_file():
        raise OnboardingError(f"导入文件不存在：{path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        if kind == "all":
            raise OnboardingError("CSV 导入必须用 --kind 指定 products、staff 或 orders")
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return {kind: [dict(row) for row in csv.DictReader(handle) if any(row.values())]}
    if suffix not in {".xlsx", ".xlsm"}:
        raise OnboardingError("只支持 UTF-8 CSV 或 XLSX 文件")

    workbook = load_workbook(path, read_only=True, data_only=True)
    targets = list(TEMPLATE_HEADERS) if kind == "all" else [kind]
    result = {}
    for target in targets:
        if target not in workbook.sheetnames:
            if kind == "all":
                continue
            raise OnboardingError(f"工作簿缺少 {target} 工作表")
        result[target] = _rows_from_sheet(workbook[target])
    if not result:
        raise OnboardingError("工作簿中没有 products、staff 或 orders 工作表")
    return result


def _value(row: dict[str, Any], *keys: str) -> Any:
    normalized = {str(key).strip().lower(): value for key, value in row.items()}
    for key in keys:
        if key.lower() in normalized:
            return normalized[key.lower()]
    return None


def _required_text(row: dict[str, Any], keys: tuple[str, ...], label: str, limit: int) -> str:
    value = str(_value(row, *keys) or "").strip()
    if not value:
        raise OnboardingError(f"{label}不能为空")
    return value[:limit]


def _money(value: Any, label: str, *, allow_zero: bool = True) -> Decimal:
    try:
        amount = Decimal(str(value).replace(",", "").replace("¥", "").replace("￥", "").strip())
    except (InvalidOperation, AttributeError) as exc:
        raise OnboardingError(f"{label}格式不正确") from exc
    if amount < 0 or (not allow_zero and amount == 0):
        raise OnboardingError(f"{label}必须大于{'等于' if allow_zero else ''} 0")
    return amount.quantize(Decimal("0.01"))


def _integer(value: Any, label: str, *, minimum: int = 0) -> int:
    try:
        number = int(float(value))
    except (TypeError, ValueError) as exc:
        raise OnboardingError(f"{label}必须是整数") from exc
    if number < minimum:
        raise OnboardingError(f"{label}不能小于 {minimum}")
    return number


def _boolean(value: Any, default: bool = False) -> bool:
    if value in (None, ""):
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "y", "on", "是", "上架"}


def _date_time(value: Any, label: str) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in (None, "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.fromisoformat(text) if fmt is None else datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise OnboardingError(f"{label}必须是 YYYY-MM-DD 或 YYYY-MM-DD HH:MM:SS")


def _role_for_staff(tenant_id: int, role_name: str) -> models.Role:
    role_name = (role_name or "客服").strip()
    if role_name not in ROLE_TEMPLATES:
        raise OnboardingError(f"员工角色必须是：{', '.join(ROLE_TEMPLATES)}")
    role = models.Role.query.filter_by(tenant_id=tenant_id, name=role_name).first()
    if role:
        return role
    template = ROLE_TEMPLATES[role_name]
    menus = models.Menu.query.filter(models.Menu.path.in_(list(template["menu_paths"]))).all()
    parent_ids = {menu.pid for menu in menus if menu.pid}
    parents = models.Menu.query.filter(models.Menu.id.in_(parent_ids)).all() if parent_ids else []
    role = models.Role(
        tenant_id=tenant_id,
        name=role_name,
        desc=template["desc"],
        perm_codes=__import__("json").dumps(template["perm_codes"], ensure_ascii=False),
        is_super=0,
    )
    role.menus = list({menu.id: menu for menu in menus + parents}.values())
    db.session.add(role)
    db.session.flush()
    return role


def _import_product(tenant_id: int, row: dict[str, Any]) -> models.Product:
    name = _required_text(row, ("name", "商品名称", "品名"), "商品名称", 128)
    price = _money(_value(row, "price", "价格", "售价"), "商品价格", allow_zero=False)
    stock_raw = _value(row, "stock", "库存", "number")
    stock = _integer(0 if stock_raw in (None, "") else stock_raw, "库存")
    if models.Product.query.filter_by(tenant_id=tenant_id, name=name).first():
        raise OnboardingError("同名商品已存在")
    category_name = str(_value(row, "category", "分类") or "").strip()[:64]
    category = None
    if category_name:
        category = models.Category.query.filter_by(tenant_id=tenant_id, name=category_name, level=1).first()
        if not category:
            category = models.Category(tenant_id=tenant_id, name=category_name, level=1)
            db.session.add(category)
            db.session.flush()
    product = models.Product(
        tenant_id=tenant_id,
        name=name,
        price=price,
        number=stock,
        introduce=str(_value(row, "description", "introduce", "商品说明", "描述") or "").strip()[:2000] or None,
        state=2 if _boolean(_value(row, "publish", "上架"), default=False) else 0,
        cid_one=category.id if category else None,
    )
    db.session.add(product)
    db.session.flush()
    return product


def _import_staff(tenant_id: int, row: dict[str, Any]) -> models.User:
    username = _required_text(row, ("username", "账号", "name"), "员工账号", 32)
    if not re.fullmatch(r"[A-Za-z0-9_.-]{3,32}", username):
        raise OnboardingError("员工账号格式不正确")
    if models.User.query.filter_by(tenant_id=tenant_id, name=username).first():
        raise OnboardingError("员工账号已存在")
    password = _validate_password(str(_value(row, "password", "密码") or ""))
    role = _role_for_staff(tenant_id, str(_value(row, "role", "角色") or "客服"))
    user = models.User(
        tenant_id=tenant_id,
        name=username,
        nick_name=str(_value(row, "nickname", "姓名", "昵称") or username).strip()[:32],
        phone=str(_value(row, "phone", "手机号") or "").strip()[:32] or None,
        email=str(_value(row, "email", "邮箱") or "").strip()[:64] or None,
        role_id=role.id,
        is_admin=0,
    )
    user.password = password
    db.session.add(user)
    db.session.flush()
    return user


def _import_order(tenant_id: int, row: dict[str, Any]) -> models.Order:
    order_no = _required_text(row, ("order_no", "订单号"), "订单号", 64)
    if models.Order.query.filter_by(tenant_id=tenant_id, order_no=order_no).first():
        raise OnboardingError("订单号已存在")
    phone = _required_text(row, ("customer_phone", "顾客手机号", "手机号"), "顾客手机号", 32)
    product_name = _required_text(row, ("product_name", "商品名称"), "商品名称", 128)
    product = models.Product.query.filter_by(tenant_id=tenant_id, name=product_name).first()
    if not product:
        raise OnboardingError("商品不存在，请先导入商品")
    quantity = _integer(_value(row, "quantity", "数量") or 1, "数量", minimum=1)
    unit_raw = _value(row, "unit_price", "单价")
    unit_price = _money(product.price if unit_raw in (None, "") else unit_raw, "单价", allow_zero=False)
    status = str(_value(row, "status", "状态") or "completed").strip().lower()
    if status not in VALID_ORDER_STATUS:
        raise OnboardingError(f"订单状态必须是：{', '.join(sorted(VALID_ORDER_STATUS))}")
    created_at = _date_time(_value(row, "created_at", "下单时间"), "下单时间") or datetime.now()
    paid_at = _date_time(_value(row, "paid_at", "付款时间"), "付款时间")
    if status in PAID_STATUS and not paid_at:
        paid_at = created_at
    customer = models.Customer.query.filter_by(tenant_id=tenant_id, phone=phone).first()
    if not customer:
        customer = models.Customer(
            tenant_id=tenant_id,
            phone=phone,
            nickname=str(_value(row, "customer_name", "顾客姓名") or "历史顾客").strip()[:64],
            status=1,
        )
        db.session.add(customer)
        db.session.flush()
    total = unit_price * quantity
    order = models.Order(
        tenant_id=tenant_id,
        order_no=order_no,
        customer_id=customer.id,
        price=total,
        goods_amount=total,
        freight_amount=0,
        number=quantity,
        status=status,
        pay_status=1 if status in PAID_STATUS else 0,
        is_send=1 if status in {"shipped", "completed"} else 0,
        address=str(_value(row, "address", "地址") or "").strip()[:255] or None,
        paid_at=paid_at,
        attribution_source="historical_import",
        attribution_medium="delivery_onboarding",
        create_time=created_at,
    )
    db.session.add(order)
    db.session.flush()
    db.session.add(models.OrderDetail(
        tenant_id=tenant_id,
        oid=order.id,
        pid=product.id,
        number=quantity,
        price=unit_price,
        total_price=total,
    ))
    if status in PAID_STATUS:
        db.session.add(models.Payment(
            tenant_id=tenant_id,
            order_id=order.id,
            out_trade_no=f"IMP-{tenant_id}-{order_no}"[:64],
            channel="historical_import",
            amount=total,
            status="success" if status != "refunded" else "refunded",
            transaction_id=f"IMPORT-{order_no}"[:64],
            paid_at=paid_at,
            raw_notify="交付初始化导入的历史订单",
            create_time=created_at,
        ))
    db.session.flush()
    return order


IMPORTERS = {"products": _import_product, "staff": _import_staff, "orders": _import_order}


def import_initial_data(tenant_slug: str, rows_by_kind: dict[str, list[dict[str, Any]]], *, dry_run: bool = False) -> dict[str, Any]:
    with tenant_context(disabled=True):
        tenant = models.Tenant.query.filter_by(slug=tenant_slug).first()
        if not tenant:
            raise OnboardingError(f"租户 {tenant_slug} 不存在")
        if (getattr(tenant, "data_mode", None) or "live") != "live":
            raise OnboardingError("正式数据导入只允许作用于 live 租户")
        if dry_run and db.session.get_bind().dialect.name == "sqlite":
            # Python sqlite3 legacy mode does not BEGIN on SELECT. Without an
            # explicit outer transaction, releasing a SAVEPOINT can persist a
            # supposed dry run before Session.rollback() gets a chance to act.
            db.session.execute(text("BEGIN"))
        ensure_base_menus()
        created: dict[str, list[dict[str, Any]]] = {kind: [] for kind in rows_by_kind}
        errors: list[dict[str, Any]] = []
        for kind in ("products", "staff", "orders"):
            if kind not in rows_by_kind:
                continue
            importer = IMPORTERS[kind]
            for row_number, row in enumerate(rows_by_kind[kind], start=2):
                try:
                    if dry_run:
                        entity = importer(tenant.id, row)
                    else:
                        with db.session.begin_nested():
                            entity = importer(tenant.id, row)
                    created[kind].append({"row": row_number, "id": entity.id, "name": getattr(entity, "name", None) or getattr(entity, "order_no", None)})
                except Exception as exc:
                    errors.append({"kind": kind, "row": row_number, "reason": str(exc), "source": dict(row)})
        if dry_run:
            db.session.rollback()
        else:
            db.session.commit()
        return {
            "tenant": tenant_slug,
            "dry_run": dry_run,
            "created": created,
            "created_count": sum(len(items) for items in created.values()),
            "error_count": len(errors),
            "errors": errors,
        }


def write_error_report(errors: list[dict[str, Any]], output_path: str | Path) -> str:
    path = Path(output_path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["kind", "row", "reason", "source"])
        writer.writeheader()
        for error in errors:
            writer.writerow({**error, "source": __import__("json").dumps(error.get("source") or {}, ensure_ascii=False)})
    return str(path)


def demo_business_counts(tenant_slug: str = "demo") -> dict[str, int]:
    with tenant_context(disabled=True):
        tenant = models.Tenant.query.filter_by(slug=tenant_slug).first()
        if not tenant:
            raise OnboardingError(f"租户 {tenant_slug} 不存在")
        if (getattr(tenant, "data_mode", None) or "live") != "demo":
            raise OnboardingError("拒绝操作：目标租户不是 demo 数据模式")
        order_ids = [row.id for row in models.Order.query.filter_by(tenant_id=tenant.id).filter(models.Order.order_no.like("DEMO-BIZ-%")).all()]
        conversation_ids = [row.id for row in models.SupportConversation.query.filter_by(tenant_id=tenant.id).filter(models.SupportConversation.session_id.like("demo-%")).all()]
        return {
            "orders": len(order_ids),
            "order_details": models.OrderDetail.query.filter_by(tenant_id=tenant.id).filter(models.OrderDetail.oid.in_(order_ids)).count() if order_ids else 0,
            "payments": models.Payment.query.filter_by(tenant_id=tenant.id).filter(models.Payment.order_id.in_(order_ids)).count() if order_ids else 0,
            "agent_traces": models.AgentTrace.query.filter_by(tenant_id=tenant.id).filter(models.AgentTrace.session_id.like("demo-%")).count(),
            "support_conversations": len(conversation_ids),
            "support_messages": models.SupportMessage.query.filter_by(tenant_id=tenant.id).filter(models.SupportMessage.conversation_id.in_(conversation_ids)).count() if conversation_ids else 0,
            "ad_campaigns": models.AdCampaignDaily.query.filter_by(tenant_id=tenant.id, source="demo").count(),
            "store_knowledge": models.StoreKnowledge.query.filter_by(tenant_id=tenant.id, source_type="demo").count(),
            "customers": models.Customer.query.filter(models.Customer.tenant_id == tenant.id, models.Customer.phone.in_(DEMO_CUSTOMER_PHONES)).count(),
        }


def purge_demo_business(tenant_slug: str = "demo", *, execute: bool = False) -> dict[str, Any]:
    counts = demo_business_counts(tenant_slug)
    if not execute:
        return {"tenant": tenant_slug, "executed": False, "counts": counts}
    with tenant_context(disabled=True):
        tenant = models.Tenant.query.filter_by(slug=tenant_slug).first()
        if not tenant or (getattr(tenant, "data_mode", None) or "live") != "demo":
            raise OnboardingError("拒绝操作：只允许清理 demo 数据模式租户")
        order_ids = [row.id for row in models.Order.query.filter_by(tenant_id=tenant.id).filter(models.Order.order_no.like("DEMO-BIZ-%")).all()]
        conversation_ids = [row.id for row in models.SupportConversation.query.filter_by(tenant_id=tenant.id).filter(models.SupportConversation.session_id.like("demo-%")).all()]
        customer_ids = [row.id for row in models.Customer.query.filter(models.Customer.tenant_id == tenant.id, models.Customer.phone.in_(DEMO_CUSTOMER_PHONES)).all()]

        if conversation_ids:
            models.SupportMessage.query.filter(models.SupportMessage.tenant_id == tenant.id, models.SupportMessage.conversation_id.in_(conversation_ids)).delete(synchronize_session=False)
        models.SupportConversation.query.filter(models.SupportConversation.tenant_id == tenant.id, models.SupportConversation.session_id.like("demo-%")).delete(synchronize_session=False)
        models.AgentConversation.query.filter(models.AgentConversation.tenant_id == tenant.id, models.AgentConversation.session_id.like("demo-%")).delete(synchronize_session=False)
        models.AgentTrace.query.filter(models.AgentTrace.tenant_id == tenant.id, models.AgentTrace.session_id.like("demo-%")).delete(synchronize_session=False)
        if order_ids:
            models.Refund.query.filter(models.Refund.tenant_id == tenant.id, models.Refund.order_id.in_(order_ids)).delete(synchronize_session=False)
            models.Payment.query.filter(models.Payment.tenant_id == tenant.id, models.Payment.order_id.in_(order_ids)).delete(synchronize_session=False)
            models.Express.query.filter(models.Express.tenant_id == tenant.id, models.Express.oid.in_(order_ids)).delete(synchronize_session=False)
            models.OrderDetail.query.filter(models.OrderDetail.tenant_id == tenant.id, models.OrderDetail.oid.in_(order_ids)).delete(synchronize_session=False)
            models.Order.query.filter(models.Order.tenant_id == tenant.id, models.Order.id.in_(order_ids)).delete(synchronize_session=False)
        models.AdCampaignDaily.query.filter_by(tenant_id=tenant.id, source="demo").delete(synchronize_session=False)
        models.StoreKnowledge.query.filter_by(tenant_id=tenant.id, source_type="demo").delete(synchronize_session=False)
        if customer_ids:
            models.CartItem.query.filter(models.CartItem.tenant_id == tenant.id, models.CartItem.customer_id.in_(customer_ids)).delete(synchronize_session=False)
            models.CustomerAddress.query.filter(models.CustomerAddress.tenant_id == tenant.id, models.CustomerAddress.customer_id.in_(customer_ids)).delete(synchronize_session=False)
            models.SmsCode.query.filter(models.SmsCode.tenant_id == tenant.id, models.SmsCode.phone.in_(DEMO_CUSTOMER_PHONES)).delete(synchronize_session=False)
            models.Customer.query.filter(models.Customer.tenant_id == tenant.id, models.Customer.id.in_(customer_ids)).delete(synchronize_session=False)
        restored = models.Product.query.filter_by(tenant_id=tenant.id, name="智能手环 S7", number=3).update({"number": 150}, synchronize_session=False)
        db.session.commit()
        return {"tenant": tenant_slug, "executed": True, "counts": counts, "stock_restored": int(restored or 0)}


def inspect_tenant(tenant_slug: str) -> dict[str, Any]:
    with tenant_context(disabled=True):
        tenant = models.Tenant.query.filter_by(slug=tenant_slug).first()
        if not tenant:
            raise OnboardingError(f"租户 {tenant_slug} 不存在")
        return {
            "tenant": tenant.to_dict(),
            "counts": {
                "staff": models.User.query.filter_by(tenant_id=tenant.id).count(),
                "products": models.Product.query.filter_by(tenant_id=tenant.id).count(),
                "orders": models.Order.query.filter_by(tenant_id=tenant.id).count(),
                "customers": models.Customer.query.filter_by(tenant_id=tenant.id).count(),
                "freight_templates": models.FreightTemplate.query.filter_by(tenant_id=tenant.id).count(),
            },
        }
